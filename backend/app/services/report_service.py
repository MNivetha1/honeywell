"""
Report Generation Service
Generates professional JSON and formatted Excel (.xlsx) validation reports.
"""

import os
import json
import io
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime
from typing import Dict, Any, List
from sqlalchemy.orm import Session
from ..database.models import ComparisonModel, ImagePairModel
from .config_service import ConfigService

class ReportService:
    REPORTS_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "../../../reports"))

    @classmethod
    def generate_json_report(cls, db: Session) -> Dict[str, Any]:
        """Generates structured JSON validation report."""
        cfg = ConfigService.get_config(db)
        comparisons = db.query(ComparisonModel).order_by(ComparisonModel.pair_id).all()
        pairs_dict = {p.pair_id: p for p in db.query(ImagePairModel).all()}

        records = []
        pass_count = 0
        fail_count = 0

        for c in comparisons:
            p = pairs_dict.get(c.pair_id)
            if c.result == "PASS":
                pass_count += 1
            else:
                fail_count += 1

            records.append({
                "pair_id": c.pair_id,
                "category_name": p.category_name if p else "Uncategorized",
                "expected_image": f"expected/{c.pair_id}.png",
                "actual_image": f"actual/{c.pair_id}.png",
                "constraints": c.constraints_applied or {
                    "text_excluded": cfg.feature_exclusion.exclude_text,
                    "icons_excluded": cfg.feature_exclusion.exclude_icons,
                    "arc_lines_excluded": cfg.feature_exclusion.exclude_arc_lines,
                    "legend_excluded": cfg.feature_exclusion.exclude_legend
                },
                "similarity_score": c.similarity_score,
                "threshold": c.threshold,
                "result": c.result,
                "margin": c.margin,
                "metrics": {
                    "iou": c.iou,
                    "pixel_agreement": c.pixel_agreement,
                    "area_similarity": c.area_similarity,
                    "intensity_agreement": c.intensity_agreement,
                    "spatial_similarity": c.spatial_similarity
                },
                "shift_dx": c.shift_dx,
                "shift_dy": c.shift_dy,
                "processing_time_ms": c.processing_time_ms,
                "ground_truth_label": p.ground_truth_label if p else None,
                "created_at": c.created_at.isoformat() if c.created_at else None
            })

        total = len(records)
        report = {
            "title": "Precipitation Rendering Functional Equivalence Validation Report",
            "generated_at": datetime.utcnow().isoformat(),
            "summary": {
                "total_pairs_evaluated": total,
                "pass_count": pass_count,
                "fail_count": fail_count,
                "pass_rate_percent": round((pass_count / max(1, total)) * 100.0, 1),
                "average_similarity_percent": round(sum(r["similarity_score"] for r in records) / max(1, total), 1),
                "threshold_percent": cfg.decision.threshold_percent,
                "weights": {
                    "overlap_weight": cfg.scoring.overlap_weight,
                    "intensity_weight": cfg.scoring.intensity_weight,
                    "area_weight": cfg.scoring.area_weight,
                    "spatial_weight": cfg.scoring.spatial_weight
                }
            },
            "results": records
        }

        # Save copy to reports/ directory
        os.makedirs(cls.REPORTS_DIR, exist_ok=True)
        timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
        report_file = os.path.join(cls.REPORTS_DIR, f"equivalence_report_{timestamp}.json")
        with open(report_file, "w", encoding="utf-8") as f:
            json.dump(report, f, indent=2)

        return report

    @classmethod
    def generate_excel_report_bytes(cls, db: Session) -> bytes:
        """Generates styled Excel (.xlsx) report with Summary and Detail sheets."""
        cfg = ConfigService.get_config(db)
        comparisons = db.query(ComparisonModel).order_by(ComparisonModel.pair_id).all()
        pairs_dict = {p.pair_id: p for p in db.query(ImagePairModel).all()}

        wb = openpyxl.Workbook()
        
        # 1. Sheet 1: Executive Summary
        ws_sum = wb.active
        ws_sum.title = "Executive Summary"
        ws_sum.views.sheetView[0].showGridLines = True

        title_font = Font(name="Segoe UI", size=16, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="0A192F", end_color="0A192F", fill_type="solid")
        section_font = Font(name="Segoe UI", size=12, bold=True, color="0A192F")
        bold_font = Font(name="Segoe UI", size=10, bold=True)
        regular_font = Font(name="Segoe UI", size=10)
        
        ws_sum.merge_cells("A1:E1")
        ws_sum["A1"] = "Precipitation Functional Equivalence Analysis Summary"
        ws_sum["A1"].font = title_font
        ws_sum["A1"].fill = header_fill
        ws_sum["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws_sum.row_dimensions[1].height = 40

        total = len(comparisons)
        pass_c = sum(1 for c in comparisons if c.result == "PASS")
        fail_c = sum(1 for c in comparisons if c.result == "FAIL")
        avg_sim = round(sum(c.similarity_score for c in comparisons) / max(1, total), 1) if total > 0 else 0.0

        summary_rows = [
            ("Report Generated At (UTC):", datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S")),
            ("Total Image Pairs Evaluated:", total),
            ("PASS Count:", pass_c),
            ("FAIL Count:", fail_c),
            ("Pass Rate (%):", f"{round((pass_c / max(1, total)) * 100.0, 1)}%"),
            ("Average Similarity (%):", f"{avg_sim}%"),
            ("PASS/FAIL Decision Threshold (%):", f"{cfg.decision.threshold_percent}%"),
            ("Overlap IoU Weight (%):", f"{cfg.scoring.overlap_weight}%"),
            ("Intensity Agreement Weight (%):", f"{cfg.scoring.intensity_weight}%"),
            ("Area Similarity Weight (%):", f"{cfg.scoring.area_weight}%"),
            ("Spatial Similarity Weight (%):", f"{cfg.scoring.spatial_weight}%"),
            ("Cosmetic Text Excluded:", "YES" if cfg.feature_exclusion.exclude_text else "NO"),
            ("Cosmetic Icons Excluded:", "YES" if cfg.feature_exclusion.exclude_icons else "NO"),
            ("Radar Arc Lines Excluded:", "YES" if cfg.feature_exclusion.exclude_arc_lines else "NO"),
            ("Radar Legend Excluded:", "YES" if cfg.feature_exclusion.exclude_legend else "NO"),
        ]

        ws_sum.cell(row=3, column=1, value="System Configuration & Results Overview").font = section_font
        for r_idx, (label, val) in enumerate(summary_rows, start=4):
            ws_sum.cell(row=r_idx, column=1, value=label).font = bold_font
            ws_sum.cell(row=r_idx, column=2, value=val).font = regular_font

        ws_sum.column_dimensions['A'].width = 36
        ws_sum.column_dimensions['B'].width = 30

        # 2. Sheet 2: Evaluation Details
        ws_det = wb.create_sheet(title="Pair Evaluation Details")
        ws_det.views.sheetView[0].showGridLines = True

        headers = [
            "Pair ID", "Variation Category", "Expected Image", "Actual Image",
            "Constraints Applied", "Similarity (%)", "Threshold (%)", "Result", "Margin (%)",
            "IoU", "Pixel Agreement", "Area Similarity", "Intensity Agreement",
            "Spatial Similarity", "dx (px)", "dy (px)", "Proc Time (ms)", "Ground Truth"
        ]

        th_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        th_fill = PatternFill(start_color="1B3B6F", end_color="1B3B6F", fill_type="solid")
        pass_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
        fail_fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
        pass_font = Font(name="Segoe UI", size=10, bold=True, color="155724")
        fail_font = Font(name="Segoe UI", size=10, bold=True, color="721C24")

        for col_num, h_text in enumerate(headers, 1):
            cell = ws_det.cell(row=1, column=col_num, value=h_text)
            cell.font = th_font
            cell.fill = th_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws_det.row_dimensions[1].height = 25

        for row_num, c in enumerate(comparisons, 2):
            p = pairs_dict.get(c.pair_id)
            c_name = p.category_name if p else "N/A"
            gt = p.ground_truth_label if p else "N/A"
            
            constraints_str = "Text, Icons, Arcs, Legend" if c.constraints_applied else "None"

            row_values = [
                c.pair_id,
                c_name,
                f"expected/{c.pair_id}.png",
                f"actual/{c.pair_id}.png",
                constraints_str,
                c.similarity_score,
                c.threshold,
                c.result,
                c.margin,
                c.iou,
                c.pixel_agreement,
                c.area_similarity,
                c.intensity_agreement,
                c.spatial_similarity,
                c.shift_dx,
                c.shift_dy,
                c.processing_time_ms,
                gt
            ]

            for col_num, val in enumerate(row_values, 1):
                cell = ws_det.cell(row=row_num, column=col_num, value=val)
                cell.font = regular_font
                cell.alignment = Alignment(horizontal="center" if col_num not in [2, 5] else "left", vertical="center")
                
                # Highlight PASS / FAIL
                if col_num == 8:  # Result column
                    if val == "PASS":
                        cell.fill = pass_fill
                        cell.font = pass_font
                    else:
                        cell.fill = fail_fill
                        cell.font = fail_font

            ws_det.row_dimensions[row_num].height = 20

        # Adjust column widths
        for col in ws_det.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws_det.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()
